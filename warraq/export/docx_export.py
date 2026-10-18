"""Word (.docx) export with proper right-to-left paragraphs."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..core.document import TITLE
from . import SIDE_WORDS


def _rtl_paragraph(p) -> None:
    ppr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def _arabic_run(p, text: str, font: str, size: float, bold: bool = False, color: RGBColor | None = None):
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rpr.append(rtl)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)
    if bold:
        bcs = OxmlElement("w:bCs")
        rpr.append(bcs)
    if color is not None:
        run.font.color.rgb = color
    return run


def _para(doc, text: str, font: str, size: float, *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=None, space_after=4, label: str | None = None, label_color=None):
    p = doc.add_paragraph()
    _rtl_paragraph(p)
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if label:
        _arabic_run(p, label + " ", font, size * 0.9, bold=True, color=label_color)
    _arabic_run(p, text, font, size, bold=bold, color=color)
    return p


def render(views, settings, doc_model, translations: list[dict] | None = None):
    out = settings.output
    fs, ms = settings.footnotes, settings.margins
    font, size = out.docx_font or "Traditional Arabic", float(out.docx_font_size or 15)
    muted = RGBColor(0x80, 0x78, 0x6C)
    accent = RGBColor(0x9A, 0x7B, 0x3F)
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size)
    sect = doc.sections[0]
    sect_pr = sect._sectPr
    bidi = OxmlElement("w:bidi")
    sect_pr.append(bidi)
    for i, v in enumerate(views):
        tr = translations[i] if translations else None
        if i > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        label = f"Page {v.number}" + (f"  ·  {v.printed_number}" if v.printed_number else "")
        hp = doc.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hp.add_run(label)
        r.font.size = Pt(9)
        r.font.color.rgb = muted
        for h in v.headers:
            _para(doc, h, font, size * 0.75, align=WD_ALIGN_PARAGRAPH.CENTER, color=muted)
        for b in v.blocks:
            if b.kind == TITLE:
                _para(doc, " ".join(b.text.split()), font, size * 1.25, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=accent, space_after=8)
                continue
            lines = b.text.split("\n")
            if settings.text.line_mode == "preserve":
                for line in lines:
                    if line.strip():
                        _para(doc, line, font, size, space_after=0)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                for para in b.text.split("\n\n"):
                    if para.strip():
                        _para(doc, " ".join(para.split()), font, size)
        if tr and tr.get("body"):
            for para in tr["body"].split("\n"):
                if para.strip():
                    p = doc.add_paragraph(para)
                    p.paragraph_format.space_after = Pt(4)
        if v.margins:
            _para(doc, ms.section_header, font, size * 0.7, align=WD_ALIGN_PARAGRAPH.RIGHT, color=muted, space_after=2)
            for m in v.margins:
                pos = f"({SIDE_WORDS.get(m.side or '', m.side)}) " if ms.include_position and m.side else ""
                _para(doc, pos + " ".join(m.text.split()), font, size * 0.8, label=m.label, label_color=accent, space_after=2)
        if v.footnotes:
            _para(doc, fs.section_header, font, size * 0.7, align=WD_ALIGN_PARAGRAPH.RIGHT, color=muted, space_after=2)
            for f in v.footnotes:
                _para(doc, f.text, font, size * 0.8, label=f.label, label_color=accent, space_after=2)
                if tr and tr.get(("footnote", f.label)):
                    p = doc.add_paragraph(f"{f.label} {tr[('footnote', f.label)]}")
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9)
    return doc


def write(path: Path, views, settings, doc_model, translations=None) -> Path:
    render(views, settings, doc_model, translations).save(str(path))
    return Path(path)
